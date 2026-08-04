import os,time
from pptx import Presentation
from pptx.util import Pt,Inches
from pptx.enum.text import PP_ALIGN
from pptx.oxml.xmlchemy import OxmlElement
from pptx.oxml.ns import qn

BASE_DIR=os.path.dirname(os.path.abspath(__file__))
FONT_DIR=os.path.join(BASE_DIR,"fonts")
DEFAULT_FONT="B Nazanin"
SLIDE_WIDTH=13.33
SLIDE_HEIGHT=7.5

def font_exists(font):
    if not os.path.exists(FONT_DIR): return False
    target=font.lower().replace("-","").replace(" ","")
    for _,_,files in os.walk(FONT_DIR):
        for f in files:
            if target in os.path.splitext(f)[0].lower().replace("-","").replace(" ",""): return True
    return False

FONT=DEFAULT_FONT if font_exists(DEFAULT_FONT) else "Arial"

def set_run_font(run,size,bold=False):
    run.font.name=FONT
    run.font.size=Pt(size)
    run.font.bold=bold
    rPr=run._r.get_or_add_rPr()
    rPr.set(qn("a:lang"),"fa-IR")
    for tag in ["latin","ea","cs"]:
        n=rPr.find(qn(f"a:{tag}"))
        if n is None:
            n=OxmlElement(f"a:{tag}")
            rPr.append(n)
        n.set("typeface",FONT)

def set_rtl(p):
    pPr=p._p.get_or_add_pPr()
    rtl=pPr.find(qn("a:rtl"))
    if rtl is None:
        rtl=OxmlElement("a:rtl")
        pPr.append(rtl)
    rtl.set(qn("a:val"),"1")
    bidi=pPr.find(qn("a:bidi"))
    if bidi is None:
        bidi=OxmlElement("a:bidi")
        pPr.append(bidi)
    bidi.set(qn("a:val"),"1")

def add_textbox(slide,text,x,y,w,h,size,bold=False,bullet=False):
    box=slide.shapes.add_textbox(Inches(x),Inches(y),Inches(w),Inches(h))
    frame=box.text_frame
    frame.clear()
    frame.word_wrap=True
    frame.margin_left=0
    frame.margin_right=0
    frame.margin_top=0
    frame.margin_bottom=0

    lines=text if isinstance(text,list) else text.split("\n")
    lines=[i for i in lines if i.strip()]

    for i,line in enumerate(lines):
        p=frame.paragraphs[0] if i==0 else frame.add_paragraph()
        set_rtl(p)
        p.text=("•  "+line) if bullet else line
        for r in p.runs:
            set_run_font(r,size,bold)

    for p in frame.paragraphs:
        p.alignment=PP_ALIGN.RIGHT
        set_rtl(p)

    return box

def add_background(slide,path):
    if path and os.path.exists(path):
        slide.shapes.add_picture(path,0,0,width=Inches(SLIDE_WIDTH),height=Inches(SLIDE_HEIGHT))

def add_image(slide,path):
    if path and os.path.exists(path):
        slide.shapes.add_picture(path,Inches(9),Inches(2),width=Inches(3))

def create_ppt(data,output):
    prs=Presentation()
    prs.slide_width=Inches(SLIDE_WIDTH)
    prs.slide_height=Inches(SLIDE_HEIGHT)

    slide=prs.slides.add_slide(prs.slide_layouts[6])
    add_textbox(slide,data.get("title",""),1,2.5,11,1,42,True)

    for item in data.get("slides",[]):
        slide=prs.slides.add_slide(prs.slide_layouts[6])
        add_background(slide,item.get("background"))

        add_textbox(slide,item.get("title",""),0.8,0.3,11.8,0.8,32,True)

        add_textbox(slide,item.get("content",[]),0.7,1.4,11.8,4.8,22,False,True)

        add_image(slide,item.get("image"))

    prs.save(output)

if __name__=="__main__":
    data={
        "title":"دنیای هوش مصنوعی",
        "slides":[
            {
                "title":"هوش مصنوعی چیست؟",
                "content":[
                    "هوش مصنوعی شاخه‌ای از علوم کامپیوتر است.",
                    "هدف آن ساخت سیستم‌های هوشمند است.",
                    "امروزه در بسیاری از صنایع استفاده می‌شود."
                ]
            },
            {
                "title":"کاربردهای AI",
                "content":[
                    "پزشکی",
                    "آموزش",
                    "تجارت",
                    "رباتیک"
                ]
            },
            {
                "title":"آینده فناوری",
                "content":[
                    "هوش مصنوعی مسیر آینده کسب‌وکارها را تغییر می‌دهد.",
                    "ترکیب انسان و ماشین فرصت‌های جدید ایجاد می‌کند."
                ]
            }
        ]
    }

    output=f"presentation_{int(time.time())}.pptx"
    print("Font:",FONT)
    create_ppt(data,output)
    print("Created:",os.path.abspath(output))
